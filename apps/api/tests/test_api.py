import csv
import io
import zipfile
from datetime import timedelta

import jwt
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.config import get_settings
from app.models import Checkin, Guest, Outbox, utcnow
from app.services.exports import qr_filename
from conftest import ADMIN_PASSWORD, PG_CODE, SCREEN_TOKEN


def upload(client, path, headers, content, filename='guests.csv'):
    return client.post(path, headers=headers, files={'file': (filename, content, 'application/octet-stream')})


def checkin(system, headers, guest_index=0, counter='COUNTER_01'):
    return system.client.post('/api/v1/pg/checkins', headers=headers, json={'guest_token': system.guests[guest_index]['token'], 'counter': counter})


def test_public_invitation_minimal_data_and_unknown_token(system):
    response = system.client.get(f'/api/v1/public/invitations/{system.guests[0]["token"]}')
    assert response.status_code == 200
    assert response.json()['guest_name'] == 'Nguyễn Văn An'
    assert response.json()['max_companions'] == 3
    assert response.json()['notes'] == ''
    assert not {'email', 'phone', 'id', 'invite_token'} & response.json().keys()
    assert response.headers['cache-control'] == 'no-store'
    assert system.client.get('/api/v1/public/invitations/unknown').status_code == 404
    event = system.client.get('/api/v1/public/event').json()
    assert event['counters'] == ['COUNTER_01', 'COUNTER_02']
    assert not {'id', 'welcome_screen_token', 'pg_access_code_hash'} & event.keys()


def test_rsvp_persists_with_outbox_and_validation(system):
    path = f'/api/v1/public/invitations/{system.guests[3]["token"]}/rsvp'
    response = system.client.put(path, json={'status': 'accepted', 'companions': 3, 'notes': 'Cần ghế gần lối đi'})
    assert response.status_code == 200
    assert response.json()['rsvp_status'] == 'accepted'
    assert response.json()['notes'] == 'Cần ghế gần lối đi'
    with system.factory() as db:
        guest = db.get(Guest, system.guests[3]['id'])
        assert guest.companions == 3 and guest.notes == 'Cần ghế gần lối đi'
        outbox = db.scalar(select(Outbox))
        assert outbox.event_type == 'rsvp' and outbox.payload == {}
        assert outbox.published_at is None
    for body in [{'status': 'accepted', 'companions': 4}, {'status': 'declined', 'companions': 1}, {'status': 'pending', 'companions': 0}, {'status': 'accepted', 'companions': True}, {'status': 'accepted', 'companions': 1, 'admin': True}]:
        assert system.client.put(path, json=body).status_code == 422
    assert system.client.put(path, json={'status': 'accepted', 'companions': 0, 'notes': 'x' * 1001}).status_code == 422
    assert system.client.put(path, json={'status': 'declined', 'companions': 0}).status_code == 200


def test_login_auth_roles_and_session_expiry(system, admin_headers, pg_headers):
    assert system.client.get('/api/v1/admin/me', headers=admin_headers).json()['email'] == 'admin@example.com'
    assert system.client.get('/api/v1/admin/guests').status_code == 401
    assert system.client.get('/api/v1/admin/guests', headers=pg_headers).status_code == 403
    assert system.client.get('/api/v1/pg/guests/search?q=An', headers=admin_headers).status_code == 403
    assert system.client.post('/api/v1/admin/login', json={'email': 'admin@example.com', 'password': 'wrong'}).status_code == 401
    assert system.client.post('/api/v1/admin/login', json={'email': 'nobody@example.com', 'password': ADMIN_PASSWORD}).status_code == 401
    assert system.client.post('/api/v1/pg/session', json={'access_code': 'wrong', 'counter': 'COUNTER_01'}).status_code == 401
    assert system.client.post('/api/v1/pg/session', json={'access_code': PG_CODE, 'counter': 'UNKNOWN'}).status_code == 422
    token = jwt.encode({'sub': '1', 'role': 'admin', 'version': 1, 'iat': utcnow() - timedelta(hours=2), 'exp': utcnow() - timedelta(hours=1), 'iss': 'thienlong-api', 'aud': 'thienlong-web'}, get_settings().jwt_secret, algorithm='HS256')
    assert system.client.get('/api/v1/admin/me', headers={'Authorization': f'Bearer {token}'}).status_code == 401


def test_revoked_admin_and_pg_sessions(system, admin_headers, pg_headers):
    from app.models import AdminUser, Event
    with system.factory() as db:
        db.get(AdminUser, 1).session_version += 1
        db.get(Event, 1).session_version += 1
        db.commit()
    assert system.client.get('/api/v1/admin/me', headers=admin_headers).status_code == 401
    assert system.client.get('/api/v1/pg/guests/search?q=An', headers=pg_headers).status_code == 401


def test_checkin_atomic_duplicate_counter_and_rsvp_lock(system, pg_headers, admin_headers):
    assert checkin(system, pg_headers, counter='COUNTER_02').status_code == 403
    first = checkin(system, pg_headers)
    assert first.status_code == 200 and first.json()['status'] == 'checked_in'
    duplicate = checkin(system, pg_headers)
    assert duplicate.status_code == 409
    assert duplicate.json()['status'] == 'already_checked_in'
    assert duplicate.json()['checked_in_at'] == first.json()['checked_in_at']
    assert duplicate.json()['counter'] == 'COUNTER_01'
    assert 'detail' not in duplicate.json()
    with system.factory() as db:
        assert db.scalar(select(func.count()).select_from(Checkin)) == 1
        assert db.scalar(select(func.count()).select_from(Outbox)) == 1
        payload = db.scalar(select(Outbox)).payload
        assert set(payload) == {'id', 'name', 'company', 'checked_in_at', 'counter'}
    path = f'/api/v1/public/invitations/{system.guests[0]["token"]}/rsvp'
    assert system.client.put(path, json={'status': 'declined', 'companions': 0}).status_code == 409
    path = f'/api/v1/admin/guests/{system.guests[0]["id"]}'
    assert system.client.patch(path, headers=admin_headers, json={'companions': 1}).status_code == 409


def test_pg_lookup_manual_search_and_checkin(system, pg_headers):
    search = system.client.get('/api/v1/pg/guests/search?q=An', headers=pg_headers)
    assert search.status_code == 200 and len(search.json()) == 1
    assert search.json()[0]['name'] == 'Nguyễn Văn An'
    assert system.client.get('/api/v1/pg/guests/search?q=Công%20ty', headers=pg_headers).json() == []
    for guest in search.json():
        assert not {'email', 'phone', 'notes', 'id', 'invitation_url'} & guest.keys()
    token = search.json()[0]['guest_token']
    resolved = system.client.get(f'/api/v1/pg/guests/by-token/{token}', headers=pg_headers)
    assert resolved.status_code == 200 and resolved.json()['guest_token'] == token
    assert system.client.post('/api/v1/pg/checkins', headers=pg_headers, json={'guest_token': token, 'counter': 'COUNTER_01'}).status_code == 200
    assert system.client.get('/api/v1/pg/guests/search?q=A', headers=pg_headers).status_code == 422
    assert system.client.get('/api/v1/pg/guests/search?q=%20%20', headers=pg_headers).status_code == 422
    assert system.client.get('/api/v1/pg/guests/search?q=%25%25', headers=pg_headers).json() == []
    assert system.client.get('/api/v1/pg/guests/by-token/invalid-token-long-enough', headers=pg_headers).status_code == 404


def test_admin_pagination_filters_sort_and_edit(system, admin_headers, pg_headers):
    checkin(system, pg_headers)
    response = system.client.get('/api/v1/admin/guests?page=2&page_size=2&sort=name', headers=admin_headers)
    assert response.status_code == 200
    assert response.json()['total'] == 4 and len(response.json()['items']) == 2
    accepted = system.client.get('/api/v1/admin/guests?rsvp_status=accepted&checkin_status=not_checked_in', headers=admin_headers).json()
    assert accepted['total'] == 1 and accepted['items'][0]['name'] == 'Trần Bình'
    assert system.client.get('/api/v1/admin/guests?search=an@example.com', headers=admin_headers).json()['total'] == 1
    assert system.client.get('/api/v1/admin/guests?sort=email', headers=admin_headers).status_code == 422
    assert system.client.get('/api/v1/admin/guests?page_size=101', headers=admin_headers).status_code == 422
    guest_id = system.guests[1]['id']
    path = f'/api/v1/admin/guests/{guest_id}'
    updated = system.client.patch(path, headers=admin_headers, json={'name': 'Trần Bình mới', 'company': 'Doanh nghiệp mới', 'email': 'BINH.NEW@example.com', 'companions': 1})
    assert updated.status_code == 200
    assert updated.json()['email'] == 'binh.new@example.com'
    assert updated.json()['notes'] == ''
    assert system.client.get(path, headers=admin_headers).json()['companions'] == 1
    assert system.client.patch(path, headers=admin_headers, json={'email': 'an@example.com'}).status_code == 409
    assert system.client.patch(path, headers=admin_headers, json={'name': None}).status_code == 422
    assert system.client.get('/api/v1/admin/guests/99999', headers=admin_headers).status_code == 404
    assert system.client.get('/api/v1/admin/event', headers=admin_headers).json()['welcome_screen_url'].endswith(SCREEN_TOKEN)


def test_admin_creates_guest_with_automatic_private_token(system, admin_headers):
    body = {'name': 'Khách bổ sung', 'company': 'Thiên Long', 'email': 'NEW.GUEST@example.com', 'phone': '0901234567'}
    response = system.client.post('/api/v1/admin/guests', headers=admin_headers, json=body)
    assert response.status_code == 201
    guest = response.json()
    assert guest['email'] == 'new.guest@example.com'
    assert guest['notes'] == ''
    assert len(guest['invite_token']) == 43
    assert guest['invitation_url'].endswith('/i/' + guest['invite_token'])
    public = system.client.get('/api/v1/public/invitations/' + guest['invite_token'])
    assert public.status_code == 200 and public.json()['guest_name'] == body['name']
    qr = system.client.get(f'/api/v1/admin/guests/{guest["id"]}/qr.png', headers=admin_headers)
    assert qr.status_code == 200 and qr.content.startswith(b'\x89PNG')
    with system.factory() as db:
        stored = db.get(Guest, guest['id'])
        assert stored.invite_token == guest['invite_token']
        assert db.scalar(select(func.count()).select_from(Outbox)) == 1
    assert system.client.post('/api/v1/admin/guests', headers=admin_headers, json=body).status_code == 409
    assert system.client.post('/api/v1/admin/guests', headers=admin_headers, json={**body, 'email': '', 'invite_token': 'chosen-by-client'}).status_code == 422
    assert system.client.post('/api/v1/admin/guests', headers=admin_headers, json={**body, 'email': '', 'notes': 'Ghi từ Admin'}).status_code == 422
    assert system.client.post('/api/v1/admin/guests', headers=admin_headers, json={**body, 'email': '', 'rsvp_status': 'pending', 'companions': 1}).status_code == 422


def test_admin_deletes_one_guest_and_then_the_complete_guest_dataset(system, admin_headers, pg_headers):
    first = system.guests[0]
    assert checkin(system, pg_headers).status_code == 200
    response = system.client.delete(f'/api/v1/admin/guests/{first["id"]}', headers=admin_headers)
    assert response.status_code == 200 and response.json() == {'deleted': 1}
    assert system.client.get(f'/api/v1/public/invitations/{first["token"]}').status_code == 404
    assert system.client.get('/api/v1/admin/guests', headers=admin_headers).json()['total'] == 3
    with system.factory() as db:
        assert db.get(Guest, first['id']) is None
        assert db.scalar(select(func.count()).select_from(Checkin)) == 0
        assert db.scalars(select(Outbox).order_by(Outbox.id.desc())).first().event_type == 'guests_changed'

    assert system.client.delete('/api/v1/admin/guests/99999', headers=admin_headers).status_code == 404
    assert system.client.delete('/api/v1/admin/guests').status_code == 401
    response = system.client.delete('/api/v1/admin/guests', headers=admin_headers)
    assert response.status_code == 200 and response.json() == {'deleted': 3}
    assert system.client.get('/api/v1/admin/guests', headers=admin_headers).json()['total'] == 0
    assert system.client.delete('/api/v1/admin/guests', headers=admin_headers).json() == {'deleted': 0}
    with system.factory() as db:
        assert db.scalar(select(func.count()).select_from(Guest)) == 0
        assert db.scalar(select(func.count()).select_from(Checkin)) == 0


def test_import_preview_atomic_validation_and_duplicate_email_skip(system, admin_headers):
    content = 'name,company,email,phone,notes\nKhách mới,Công ty M,moi@example.com,0900000000,Đón tại sảnh A\nKhách lỗi,Công ty B,invalid-email,,\n'.encode()
    preview = upload(system.client, '/api/v1/admin/guests/import/preview', admin_headers, content)
    assert preview.status_code == 200
    assert preview.json()['total'] == 2 and preview.json()['valid_count'] == 1
    assert preview.json()['rows'][0]['notes'] == 'Đón tại sảnh A'
    assert preview.json()['errors'][0]['row_number'] == 3
    assert upload(system.client, '/api/v1/admin/guests/import', admin_headers, content).status_code == 422
    with system.factory() as db:
        assert db.scalar(select(func.count()).select_from(Guest)) == 4
        assert db.scalar(select(func.count()).select_from(Outbox)) == 0
    valid = 'name,company,email,phone,notes\nTrùng email,A,AN@example.com,,\nCùng tên,A,new@example.com,,Ghi chú import\nCùng tên,B,NEW@example.com,,\nCùng tên,C,,,\nCùng tên,D,,,\n'.encode()
    result = upload(system.client, '/api/v1/admin/guests/import', admin_headers, valid)
    assert result.status_code == 200 and result.json() == {'imported': 3, 'skipped': 2}
    with system.factory() as db:
        guests = db.scalars(select(Guest)).all()
        assert len(guests) == 7 and len({guest.invite_token for guest in guests}) == 7
        assert all(len(guest.invite_token) >= 32 for guest in guests)
        assert any(guest.notes == 'Ghi chú import' for guest in guests)
        assert db.scalar(select(Outbox)).event_type == 'guests_changed'


def test_import_minimal_name_only_csv_and_xlsx(system, admin_headers):
    preview = upload(system.client, '/api/v1/admin/guests/import/preview', admin_headers, 'name\nNguyễn Văn Mới\n'.encode())
    assert preview.status_code == 200 and preview.json()['valid_count'] == 1
    assert preview.json()['rows'][0]['email'] is None
    workbook = Workbook()
    workbook.active.append(['name', 'company', 'email', 'phone'])
    workbook.active.append(['Khách Excel', 'Thiên Long', 'excel@example.com', '0901234567'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = upload(system.client, '/api/v1/admin/guests/import', admin_headers, buffer.getvalue(), 'guests.xlsx')
    assert result.status_code == 200 and result.json()['imported'] == 1


@pytest.mark.parametrize('content,filename', [(b'', 'empty.csv'), (b'name\n', 'headers.csv'), (b'company\nAcme\n', 'missing.csv'), (b'name,name\nA,B\n', 'duplicate.csv'), (b'not-an-excel', 'broken.xlsx'), (b'hello', 'wrong.pdf')])
def test_import_rejects_bad_files(system, admin_headers, content, filename):
    assert upload(system.client, '/api/v1/admin/guests/import/preview', admin_headers, content, filename).status_code == 422


def test_import_formula_and_size_rejected(system, admin_headers):
    preview = upload(system.client, '/api/v1/admin/guests/import/preview', admin_headers, b'name,company\n=1+1,Acme\n')
    assert preview.status_code == 200 and preview.json()['valid_count'] == 0
    assert preview.json()['errors']
    assert upload(system.client, '/api/v1/admin/guests/import', admin_headers, b'x' * (5 * 1024 * 1024 + 1)).status_code == 413


def test_dashboard_precise_kpis_welcome_and_trend(system, pg_headers, admin_headers):
    assert system.client.get(f'/api/v1/public/welcome/{SCREEN_TOKEN}').json()['latest_guest'] is None
    checkin(system, pg_headers, 0)
    checkin(system, pg_headers, 3)
    summary = system.client.get('/api/v1/admin/dashboard/summary', headers=admin_headers).json()
    assert {key: summary[key] for key in ('total_guests', 'accepted', 'declined', 'pending', 'expected_attendance', 'checked_in', 'not_arrived', 'no_show', 'checkin_rate', 'registered_arrived')} == {'total_guests': 4, 'accepted': 2, 'declined': 1, 'pending': 1, 'expected_attendance': 4, 'checked_in': 2, 'not_arrived': 2, 'no_show': 1, 'checkin_rate': 50.0, 'registered_arrived': 4}
    assert summary['checkins_by_counter'] == [{'counter': 'COUNTER_01', 'count': 2}, {'counter': 'COUNTER_02', 'count': 0}]
    latest = system.client.get(f'/api/v1/public/welcome/{SCREEN_TOKEN}').json()['latest_guest']
    assert latest['name'] == 'Phạm Dũng'
    assert set(latest) == {'id', 'name', 'company', 'checked_in_at'}
    assert system.client.get('/api/v1/public/welcome/invalid').status_code == 404
    trend = system.client.get('/api/v1/admin/dashboard/checkins', headers=admin_headers).json()
    assert sum(row['count'] for row in trend) == 2
    assert system.client.get('/api/v1/admin/dashboard/checkins?from=2026-01-01', headers=admin_headers).status_code == 422
    assert system.client.get('/api/v1/admin/dashboard/checkins?from=2027-01-01T00:00:00Z&to=2026-01-01T00:00:00Z', headers=admin_headers).status_code == 422


def test_exports_include_not_arrived_and_neutralize_formula(system, admin_headers, pg_headers):
    with system.factory() as db:
        db.get(Guest, system.guests[1]['id']).company = '=HYPERLINK("https://example.com")'
        db.get(Guest, system.guests[1]['id']).notes = 'Cần hỗ trợ đón tiếp'
        db.commit()
    checkin(system, pg_headers)
    response = system.client.get('/api/v1/admin/export/checkins.csv', headers=admin_headers)
    assert response.status_code == 200 and response.content.startswith(b'\xef\xbb\xbf')
    rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(rows) == 5
    assert rows[0][-1] == 'Notes' and rows[2][-1] == 'Cần hỗ trợ đón tiếp'
    assert rows[2][2].startswith("'=")
    assert rows[1][6] == 'checked_in' and rows[2][6] == 'not_checked_in'
    response = system.client.get('/api/v1/admin/export/checkins.xlsx', headers=admin_headers)
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ['Guests', 'Summary']
    assert workbook['Guests'].max_row == 5
    assert workbook['Guests']['C3'].data_type != 'f'
    assert workbook['Summary'].max_row == 11
    assert system.client.get('/api/v1/admin/export/checkins.csv').status_code == 401


def test_qr_zip_mapping_and_template(system, admin_headers, pg_headers):
    response = system.client.get(f'/api/v1/admin/guests/{system.guests[0]["id"]}/qr.png', headers=admin_headers)
    assert response.status_code == 200 and response.content.startswith(b'\x89PNG\r\n\x1a\n')
    assert response.headers['content-disposition'] == 'attachment; filename="nguyen-van-an-thienlong.png"'
    response = system.client.get('/api/v1/admin/guests/qr.zip', headers=admin_headers)
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert len(archive.namelist()) == 5
        rows = list(csv.reader(io.StringIO(archive.read('mapping.csv').decode('utf-8-sig'))))
        assert len(rows) == 5
        assert 'qr/nguyen-van-an-thienlong.png' in archive.namelist()
        for row in rows[1:]:
            assert row[2].startswith('http://localhost:3000/i/invitation-token-')
            assert archive.read(row[3]).startswith(b'\x89PNG')
    assert system.client.get('/api/v1/admin/guests/qr.zip', headers=pg_headers).status_code == 403
    template = system.client.get('/api/v1/admin/guests/template.csv', headers=admin_headers)
    assert template.status_code == 200
    assert template.content.decode('utf-8-sig').splitlines()[0] == 'name,company,email,phone,notes'


def test_qr_filename_removes_vietnamese_accents_and_uses_brand_suffix():
    assert qr_filename(Guest(id=7, name='Đặng Như Phước')) == 'dang-nhu-phuoc-thienlong.png'
    assert qr_filename(Guest(id=8, name='Diệp Gia Luật')) == 'diep-gia-luat-thienlong.png'
    assert qr_filename(Guest(id=9, name='Diệp Gia Luật'), 2) == 'diep-gia-luat-2-thienlong.png'


def test_rate_limit_and_validation_do_not_echo_secret(system):
    secret = 'do-not-echo-this-secret'
    response = system.client.post('/api/v1/admin/login', json={'email': 'invalid', 'password': secret})
    assert response.status_code == 422 and secret not in response.text
    for _ in range(10):
        assert system.client.post('/api/v1/admin/login', json={'email': 'admin@example.com', 'password': 'wrong'}).status_code == 401
    response = system.client.post('/api/v1/admin/login', json={'email': 'admin@example.com', 'password': ADMIN_PASSWORD})
    assert response.status_code == 429 and int(response.headers['retry-after']) > 0


def test_health_openapi_and_cors(system):
    assert system.client.get('/health').json() == {'status': 'ok'}
    assert system.client.get('/health/ready').status_code == 200
    document = system.client.get('/openapi.json').json()
    assert '/api/v1/admin/stream' in document['paths']
    assert 'CheckinResponse' in document['components']['schemas']
    good = system.client.options('/api/v1/admin/me', headers={'Origin': 'http://localhost:3000', 'Access-Control-Request-Method': 'GET', 'Access-Control-Request-Headers': 'authorization,last-event-id'})
    assert good.status_code == 200 and good.headers['access-control-allow-origin'] == 'http://localhost:3000'
    delete_preflight = system.client.options('/api/v1/admin/guests', headers={'Origin': 'http://localhost:3000', 'Access-Control-Request-Method': 'DELETE', 'Access-Control-Request-Headers': 'authorization'})
    assert delete_preflight.status_code == 200
    assert 'DELETE' in delete_preflight.headers['access-control-allow-methods']
    bad = system.client.options('/api/v1/admin/me', headers={'Origin': 'https://evil.example', 'Access-Control-Request-Method': 'GET'})
    assert bad.status_code == 400 and 'access-control-allow-origin' not in bad.headers
