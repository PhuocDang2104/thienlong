import csv
import io
import zipfile

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import Event, Guest
from app.schemas import ImportPreview, ImportRow
from app.services.guests import add_outbox, get_event

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


async def read_upload(file: UploadFile) -> tuple[bytes, str]:
    data = await file.read(MAX_BYTES + 1)
    await file.close()
    if len(data) > MAX_BYTES:
        raise HTTPException(413, 'Tệp tối đa 5 MB.')
    if not data:
        raise HTTPException(422, 'Tệp trống.')
    suffix = (file.filename or '').lower().rsplit('.', 1)[-1]
    if suffix not in ('csv', 'xlsx'):
        raise HTTPException(422, 'Chỉ hỗ trợ CSV UTF-8 và XLSX.')
    return data, suffix


def preview_import(data: bytes, suffix: str) -> ImportPreview:
    try:
        if suffix == 'csv':
            text = data.decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel  # A name-only CSV is valid and has no delimiter.
            source = csv.reader(io.StringIO(text), dialect)
            matrix = []
            for row in source:
                matrix.append(row)
                if len(matrix) > MAX_ROWS + 1:
                    raise HTTPException(422, 'Tối đa 5.000 dòng trong một lần nhập.')
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise HTTPException(413, 'Tệp Excel giải nén quá lớn.')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                matrix = []
                for row in workbook.active.iter_rows():
                    matrix.append([str(cell.value).strip() if cell.value is not None else '' for cell in row])
                    if len(matrix) > MAX_ROWS + 1:
                        raise HTTPException(422, 'Tối đa 5.000 dòng trong một lần nhập.')
            finally:
                workbook.close()
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'Không đọc được tệp. Sử dụng CSV UTF-8 hoặc Excel XLSX hợp lệ.') from None
    if not matrix:
        raise HTTPException(422, 'Tệp không có dữ liệu.')
    headers = [str(value).strip().lower() for value in matrix[0]]
    if 'name' not in headers:
        raise HTTPException(422, 'Thiếu cột bắt buộc: name. Các cột hỗ trợ: name, company, email, phone, notes.')
    nonblank_headers = [header for header in headers if header]
    if len(nonblank_headers) != len(set(nonblank_headers)):
        raise HTTPException(422, 'Tên cột bị trùng.')
    rows, errors = [], []
    total = 0
    for row_number, values in enumerate(matrix[1:], start=2):
        if not any(str(value).strip() for value in values):
            continue
        total += 1
        if len(values) > len(headers) and any(values[len(headers):]):
            errors.append({'row_number': row_number, 'message': 'Dòng có nhiều ô hơn số cột tiêu đề.'})
            continue
        record = {column: str(values[index]).strip() if index < len(values) else '' for index, column in enumerate(headers) if column in ('name', 'company', 'email', 'phone', 'notes')}
        record['email'] = record.get('email', '').lower() or None
        if any(isinstance(value, str) and value.startswith('=') for value in record.values()):
            errors.append({'row_number': row_number, 'message': 'Không chấp nhận công thức trong dữ liệu khách mời.'})
            continue
        try:
            rows.append(ImportRow(row_number=row_number, **record))
        except ValidationError as exc:
            fields = ', '.join(str(error['loc'][0]) for error in exc.errors())
            errors.append({'row_number': row_number, 'message': f'Dữ liệu không hợp lệ tại cột: {fields}.'})
    if not total:
        raise HTTPException(422, 'Tệp chưa có khách mời.')
    return ImportPreview(rows=rows, errors=errors, total=total, valid_count=len(rows))


def commit_import(db: Session, preview: ImportPreview) -> dict:
    if preview.errors:
        raise HTTPException(422, 'Tệp có dòng không hợp lệ. Sửa tất cả lỗi trong bản xem trước rồi nhập lại.')
    get_event(db)
    # Serialize import batches for stable deduplication, including concurrent uploads.
    db.scalar(select(Event).where(Event.id == 1).with_for_update())
    known = set(db.scalars(select(Guest.email).where(Guest.event_id == 1, Guest.email.is_not(None))).all())
    imported = skipped = 0
    for row in preview.rows:
        email = str(row.email).lower() if row.email else None
        if email and email in known:
            skipped += 1
            continue
        db.add(Guest(event_id=1, name=row.name, company=row.company, email=email, phone=row.phone, notes=row.notes))
        if email:
            known.add(email)
        imported += 1
    if imported:
        add_outbox(db, 'guests_changed')
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, 'Danh sách vừa thay đổi. Vui lòng xem trước và nhập lại.') from None
    return {'imported': imported, 'skipped': skipped}
