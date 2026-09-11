import csv
import io
import zipfile
from zoneinfo import ZoneInfo

import qrcode
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.models import Guest
from app.services.guests import invitation_url


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def csv_bytes(rows: list[list]) -> bytes:
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerows([[safe_cell(value) for value in row] for row in rows])
    return stream.getvalue().encode('utf-8-sig')


def qr_png(guest: Guest) -> bytes:
    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=10, border=4)
    qr.add_data(invitation_url(guest))
    qr.make(fit=True)
    image = qr.make_image(fill_color='black', back_color='white')
    stream = io.BytesIO()
    image.save(stream, format='PNG')
    return stream.getvalue()


def qr_archive(guests: list[Guest]) -> bytes:
    stream = io.BytesIO()
    mapping = [['Guest Name', 'Company', 'Invitation URL', 'QR File']]
    with zipfile.ZipFile(stream, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for guest in guests:
            filename = f'qr/guest-{guest.id:05d}.png'
            archive.writestr(filename, qr_png(guest))
            mapping.append([guest.name, guest.company, invitation_url(guest), filename])
        archive.writestr('mapping.csv', csv_bytes(mapping))
    return stream.getvalue()


def report_rows(guests: list[Guest]) -> list[list]:
    rows = [['Guest Name', 'Email', 'Company', 'Phone', 'RSVP Status', 'Companions', 'Check-in Status', 'Check-in Time (Asia/Ho_Chi_Minh)', 'Check-in Counter', 'Notes']]
    for guest in guests:
        rows.append([guest.name, guest.email or '', guest.company, guest.phone, guest.rsvp_status, guest.companions, 'checked_in' if guest.checkin else 'not_checked_in', guest.checkin.checked_in_at.astimezone(ZoneInfo('Asia/Ho_Chi_Minh')).isoformat() if guest.checkin else '', guest.checkin.counter if guest.checkin else '', guest.notes])
    return rows


def xlsx_report(guests: list[Guest], summary: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Guests'
    for row in report_rows(guests):
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='173F35')
    for column, width in {'A': 30, 'B': 32, 'C': 30, 'D': 20, 'E': 18, 'F': 14, 'G': 22, 'H': 32, 'I': 22, 'J': 44}.items():
        sheet.column_dimensions[column].width = width
    kpis = workbook.create_sheet('Summary')
    kpis.append(['Metric', 'Value', 'Definition'])
    definitions = {'total_guests': 'All invitations', 'accepted': 'RSVP accepted invitations', 'declined': 'RSVP declined invitations', 'pending': 'RSVP pending invitations', 'expected_attendance': 'Accepted guests plus their registered companions', 'checked_in': 'Invitations checked in (not actual individual headcount)', 'not_arrived': 'All invitations minus checked-in invitations', 'no_show': 'Accepted invitations not yet checked in', 'checkin_rate': 'Checked-in invitations / all invitations * 100 (%)', 'registered_arrived': 'Sum of 1 + registered companions for checked-in guests; estimated headcount'}
    for key, definition in definitions.items():
        kpis.append([key, summary[key], definition])
    kpis.column_dimensions['A'].width = 26
    kpis.column_dimensions['B'].width = 14
    kpis.column_dimensions['C'].width = 100
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
