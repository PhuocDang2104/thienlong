"""Verify admin search/filter/edit/import/export and SSE refresh on local data."""
import io
import uuid
import zipfile
from pathlib import Path

import httpx
from dotenv import dotenv_values
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]
CONFIG = dotenv_values(ROOT / 'apps/api/.env')
API = 'http://localhost:8000/api/v1'
WEB = 'http://localhost:3000'


def main():
    if CONFIG.get('APP_ENV') != 'development':
        raise RuntimeError('Admin browser checks require development configuration')
    batch = uuid.uuid4().hex[:7]
    (ROOT / '.local/screenshots').mkdir(parents=True, exist_ok=True)
    name = f'Khách quản trị {batch}'
    manual_name = f'Khách thêm tay {batch}'
    with httpx.Client(timeout=30) as client:
        login = client.post(API + '/admin/login', json={'email': CONFIG['ADMIN_EMAIL'], 'password': CONFIG['ADMIN_PASSWORD']})
        login.raise_for_status()
        auth = {'Authorization': 'Bearer ' + login.json()['access_token']}
        data = f'name,company,email,phone\n{name},Đối tác kiểm thử,admin-{batch}@example.com,0901111222\n'
        client.post(API + '/admin/guests/import', headers=auth, files={'file': ('admin.csv', data.encode(), 'text/csv')}).raise_for_status()
        guest = client.get(API + '/admin/guests', headers=auth, params={'search': name}).json()['items'][0]
        event = client.get(API + '/public/event').json()
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={'width': 1440, 'height': 1000})
            page.goto(WEB + '/admin/login')
            page.get_by_label('Email', exact=True).fill(CONFIG['ADMIN_EMAIL'])
            page.get_by_label('Mật khẩu', exact=True).fill(CONFIG['ADMIN_PASSWORD'])
            page.get_by_role('button', name='Đăng nhập', exact=True).click()
            expect(page.get_by_role('navigation', name='Điều hướng quản trị')).to_be_visible(timeout=15000)
            nav = page.get_by_role('navigation', name='Điều hướng quản trị')
            nav.get_by_role('link', name='Khách mời', exact=True).click()
            page.get_by_role('button', name='Thêm khách', exact=True).click()
            create_dialog = page.get_by_role('dialog')
            create_dialog.get_by_label('Họ và tên').fill(manual_name)
            create_dialog.get_by_label('Đơn vị / Công ty').fill('Khách phát sinh')
            create_dialog.get_by_label('Email', exact=True).fill(f'manual-{batch}@example.com')
            create_dialog.get_by_role('button', name='Thêm khách mời', exact=True).click()
            expect(create_dialog).not_to_be_visible()
            expect(page.get_by_text('Đã thêm khách và tạo link/QR riêng.', exact=True)).to_be_visible()
            created = client.get(API + '/admin/guests', headers=auth, params={'search': manual_name}).json()['items'][0]
            assert len(created['invite_token']) == 43 and created['invitation_url'].endswith('/i/' + created['invite_token'])
            assert created['notes'] == ''
            page.get_by_label('Tìm khách mời', exact=True).fill(name)
            table = page.get_by_role('table')
            expect(table.get_by_role('row')).to_have_count(2)
            table.get_by_role('button', name=name, exact=True).click()
            dialog = page.get_by_role('dialog')
            expect(dialog).to_be_visible()
            dialog.get_by_label('Đơn vị / Công ty').fill('Đơn vị đã cập nhật')
            dialog.get_by_label('Phản hồi tham dự').select_option('accepted')
            dialog.get_by_label('Người đi cùng', exact=True).fill('1')
            dialog.get_by_role('button', name='Lưu thay đổi').click()
            expect(dialog).not_to_be_visible()
            expect(table.get_by_text('Đơn vị đã cập nhật', exact=True)).to_be_visible()
            page.get_by_label('Lọc phản hồi').select_option('declined')
            expect(page.get_by_text('Không có khách mời phù hợp', exact=True)).to_be_visible()
            page.get_by_label('Lọc phản hồi').select_option('accepted')
            expect(table.get_by_role('row')).to_have_count(2)
            page.get_by_label('Số khách mỗi trang').select_option('50')

            page.screenshot(path=str(ROOT / '.local/screenshots/guests-invitation-tab.png'), full_page=True)
            page.get_by_role('tab', name='Danh sách check-in', exact=True).click()
            expect(table.get_by_role('row')).to_have_count(2)
            expect(table.get_by_text(name, exact=True)).to_be_visible()
            page.screenshot(path=str(ROOT / '.local/screenshots/guests-checkin-tab.png'), full_page=True)
            page.get_by_label('Lọc check-in').select_option('not_checked_in')
            expect(table.get_by_text(name, exact=True)).to_be_visible()
            page.get_by_label('Lọc check-in').select_option('')

            pg = client.post(API + '/pg/session', json={'access_code': CONFIG['PG_ACCESS_CODE'], 'counter': event['counters'][0]})
            pg.raise_for_status()
            response = client.post(API + '/pg/checkins', headers={'Authorization': 'Bearer ' + pg.json()['access_token']}, json={'guest_token': guest['invite_token'], 'counter': event['counters'][0]})
            response.raise_for_status()
            expect(table.get_by_text('Đã check-in', exact=True)).to_be_visible(timeout=15000)
            table.get_by_role('button', name=name, exact=True).click()
            expect(dialog.get_by_text('Khách đã check-in.', exact=False)).to_be_visible()
            dialog.get_by_label('Email', exact=True).fill(f'updated-{batch}@example.com')
            dialog.get_by_role('button', name='Lưu thay đổi').click()
            expect(dialog).not_to_be_visible()
            page.get_by_role('tab', name='Thư mời', exact=True).click()
            expect(table.get_by_text(f'updated-{batch}@example.com', exact=True).last).to_be_visible()

            nav.get_by_role('link', name='Nhập danh sách', exact=True).click()
            bad = b'name,company,email\n,Test,invalid\n'
            page.locator('#import-file').set_input_files({'name': 'invalid.csv', 'mimeType': 'text/csv', 'buffer': bad})
            page.get_by_role('button', name='Kiểm tra dữ liệu', exact=True).click()
            expect(page.get_by_role('button', name='Xác nhận nhập danh sách')).to_be_disabled()
            good = f'name,company,email,phone\nNhập qua giao diện {batch},Thiên Long,ui-{batch}@example.com,\n'.encode()
            page.locator('#import-file').set_input_files({'name': 'valid.csv', 'mimeType': 'text/csv', 'buffer': good})
            page.get_by_role('button', name='Kiểm tra dữ liệu', exact=True).click()
            confirm = page.get_by_role('button', name='Xác nhận nhập danh sách')
            expect(confirm).to_be_enabled()
            confirm.click()
            expect(page.get_by_text('1 khách mời', exact=True)).to_be_visible()

            nav.get_by_role('link', name='Xuất dữ liệu', exact=True).click()
            with page.expect_download() as pending:
                page.get_by_role('button', name='Tải báo cáo .xlsx', exact=True).click()
            workbook = load_workbook(io.BytesIO(Path(pending.value.path()).read_bytes()), read_only=True)
            assert 'Guests' in workbook.sheetnames and 'Summary' in workbook.sheetnames
            workbook.close()
            with page.expect_download(timeout=60000) as pending:
                page.get_by_role('button', name='Tải bộ QR .zip', exact=True).click()
            with zipfile.ZipFile(pending.value.path()) as archive:
                assert 'mapping.csv' in archive.namelist()
            page.screenshot(path=str(ROOT / '.local/screenshots/export-desktop.png'), full_page=True)
            browser.close()
    print('PASS: invitation/check-in tabs, RSVP appears immediately, UI add/edit guest, SSE refresh, import validation/commit, XLSX and QR ZIP downloads.')


if __name__ == '__main__':
    main()
